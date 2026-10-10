"""Rendu Word (.docx) du CV et de la lettre de motivation.

Miroir de `pdf_export.py` (mêmes données `CVContent`, même style : accent
teal, police Calibri, en-tête centré, titres de section soulignés) pour
que chaque candidature dispose des deux formats, comme le veut
`application_folder.save_application_folder`. Les deux modules restent
volontairement séparés (reportlab vs python-docx ont des APIs trop
différentes pour un seul code de mise en page partagé) mais doivent être
modifiés ensemble si le style de référence change.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from aelyn_career.models import CVContent

_ACCENT = RGBColor(0x00, 0x86, 0x8C)
_MUTED = RGBColor(0x59, 0x59, 0x59)
_INK = RGBColor(0x1A, 0x1A, 0x1A)
_FONT = "Calibri"


def _set_margins(
    document: Document, *, top_bottom_pt: float, left_right_pt: float
) -> None:
    section = document.sections[0]
    section.top_margin = Pt(top_bottom_pt)
    section.bottom_margin = Pt(top_bottom_pt)
    section.left_margin = Pt(left_right_pt)
    section.right_margin = Pt(left_right_pt)


def _run(
    paragraph, text: str, *, size: float, bold: bool = False, italic: bool = False,
    color=None,
):
    run = paragraph.add_run(text)
    run.font.name = _FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def _bottom_border(paragraph, *, color_hex: str, size_eighths_pt: int) -> None:
    """Bordure inférieure de paragraphe (équivalent du `HRFlowable`/
    `pBdr` côté reportlab/docx-js) : pas d'API python-docx de haut niveau
    pour ça, seule la manipulation OOXML directe l'expose."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size_eighths_pt))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color_hex)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_hyperlink(
    paragraph, url: str, text: str, *, size: float, color: RGBColor
) -> None:
    """python-docx n'expose aucune méthode native pour un hyperlien
    externe (contrairement à docx-js `ExternalHyperlink`) : il faut
    déclarer la relation externe sur la part du document, puis construire
    à la main l'élément `w:hyperlink` qui y pointe. Recette standard,
    jamais changée depuis les premières versions de python-docx."""
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run_element = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")

    font_element = OxmlElement("w:rFonts")
    font_element.set(qn("w:ascii"), _FONT)
    font_element.set(qn("w:hAnsi"), _FONT)
    run_props.append(font_element)

    size_element = OxmlElement("w:sz")
    size_element.set(qn("w:val"), str(int(size * 2)))
    run_props.append(size_element)

    color_element = OxmlElement("w:color")
    color_element.set(qn("w:val"), str(color))
    run_props.append(color_element)

    underline_element = OxmlElement("w:u")
    underline_element.set(qn("w:val"), "single")
    run_props.append(underline_element)

    run_element.append(run_props)
    text_element = OxmlElement("w:t")
    text_element.set(qn("xml:space"), "preserve")
    text_element.text = text
    run_element.append(text_element)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)


def _section_heading(document: Document, title: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    _run(p, title.upper(), size=10.5, bold=True, color=_ACCENT)
    _bottom_border(p, color_hex="00868C", size_eighths_pt=8)


def _entry_heading(document: Document, title_line: str, dates: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    _run(p, title_line, size=9.5, bold=True, color=_INK)
    _run(p, "   ", size=9.5)
    _run(p, dates, size=8.5, italic=True, color=_MUTED)


def _bullet(document: Document, text: str) -> None:
    p = document.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Pt(11)
    for run in list(p.runs):
        run.text = ""
    _run(p, text, size=8.5, color=_INK)


def _body(document: Document, text: str) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    _run(p, text, size=8.5, color=_INK)


def cv_to_docx_bytes(cv: CVContent, *, header_lines: list[str]) -> bytes:
    """Rend un `CVContent` en .docx une page, même style que
    `pdf_export.cv_to_pdf_bytes` (teal, Calibri, en-tête centré)."""
    document = Document()
    _set_margins(document, top_bottom_pt=21, left_right_pt=36)

    if header_lines:
        name_p = document.add_paragraph()
        name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        name_p.paragraph_format.space_after = Pt(1)
        _run(name_p, header_lines[0], size=19, bold=True, color=_ACCENT)

        if len(header_lines) > 1:
            contact_p = document.add_paragraph()
            contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            contact_p.paragraph_format.space_after = Pt(5)
            _run(contact_p, " · ".join(header_lines[1:]), size=8.3, color=_MUTED)
            _bottom_border(contact_p, color_hex="00868C", size_eighths_pt=10)

    _section_heading(document, "Profil")
    _body(document, cv.profil)

    _section_heading(document, "Expériences")
    for e in cv.experiences:
        _entry_heading(document, f"{e.role} · {e.entreprise}", e.periode)
        for puce in e.puces:
            _bullet(document, puce)

    if cv.projets:
        _section_heading(document, "Projets")
        for p in cv.projets:
            para = document.add_paragraph()
            para.paragraph_format.space_after = Pt(3)
            _run(para, p.titre, size=8.5, bold=True, color=_INK)
            _run(para, f" : {p.description}", size=8.5, color=_INK)

    _section_heading(document, "Compétences")
    for categorie, items in cv.competences.items():
        para = document.add_paragraph()
        para.paragraph_format.space_after = Pt(3)
        _run(para, f"{categorie} : ", size=8.5, bold=True, color=_INK)
        _run(para, ", ".join(items), size=8.5, color=_INK)

    _section_heading(document, "Formation")
    for f in cv.formation:
        _body(document, f)

    if cv.certifications:
        _section_heading(document, "Certifications")
        for c in cv.certifications:
            para = document.add_paragraph()
            para.paragraph_format.space_after = Pt(3)
            url = cv.certification_urls.get(c)
            if url:
                _add_hyperlink(para, url, c, size=8.5, color=_ACCENT)
            else:
                _run(para, c, size=8.5, color=_INK)

    if cv.langues:
        _section_heading(document, "Langues")
        _body(document, " · ".join(cv.langues))

    if cv.centres_interet:
        _section_heading(document, "Centres d'intérêt")
        _body(document, " · ".join(cv.centres_interet))

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def cover_letter_to_docx_bytes(text: str) -> bytes:
    """Rend une lettre de motivation (texte brut, paragraphes séparés par
    des lignes vides) en .docx, même police que le CV. `text` contient
    déjà l'en-tête (nom/coordonnées de l'expéditeur), comme pour
    `pdf_export.cover_letter_to_pdf_bytes`."""
    document = Document()
    _set_margins(document, top_bottom_pt=28, left_right_pt=34)

    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(10)
        lines = block.split("\n")
        for i, line in enumerate(lines):
            if i > 0:
                p.add_run().add_break()
            _run(p, line, size=10.5, color=_INK)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
